import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open("data/vendor/wgloop/equipment.json", encoding="utf-8") as f:
    equipment = json.load(f)

def n(v):
    return v if isinstance(v, (int, float)) else 0

rows = []
for it in equipment:
    rows.append({
        "id": it["id"],
        "name": it["name"],
        "version": it.get("version") or "",
        "slot": it.get("slot") or "",
        "category": it.get("category") or "",
        "speed": n(it.get("speed")),
        "isTwoHanded": it.get("isTwoHanded", False),
        "str": n((it.get("bonuses") or {}).get("str")),
        "rangedStr": n((it.get("bonuses") or {}).get("ranged_str")),
        "magicStr": n((it.get("bonuses") or {}).get("magic_str")),
        "prayer": n((it.get("bonuses") or {}).get("prayer")),
        "attackStab": n((it.get("offensive") or {}).get("stab")),
        "attackSlash": n((it.get("offensive") or {}).get("slash")),
        "attackCrush": n((it.get("offensive") or {}).get("crush")),
        "attackMagic": n((it.get("offensive") or {}).get("magic")),
        "attackRanged": n((it.get("offensive") or {}).get("ranged")),
        "defStab": n((it.get("defensive") or {}).get("stab")),
        "defSlash": n((it.get("defensive") or {}).get("slash")),
        "defCrush": n((it.get("defensive") or {}).get("crush")),
        "defMagic": n((it.get("defensive") or {}).get("magic")),
        "defRanged": n((it.get("defensive") or {}).get("ranged")),
    })

headers = [
    ("ID", 8),
    ("Name", 30),
    ("Version", 16),
    ("Slot", 10),
    ("Category", 18),
    ("Speed", 8),
    ("2H?", 6),
    ("Melee Str", 10),
    ("Ranged Str", 11),
    ("Magic Dmg×10", 13),
    ("Prayer", 8),
    ("Atk Stab", 10),
    ("Atk Slash", 10),
    ("Atk Crush", 10),
    ("Atk Magic", 10),
    ("Atk Ranged", 11),
    ("Def Stab", 9),
    ("Def Slash", 9),
    ("Def Crush", 9),
    ("Def Magic", 9),
    ("Def Ranged", 10),
]

fields = [
    "id","name","version","slot","category","speed","isTwoHanded",
    "str","rangedStr","magicStr","prayer",
    "attackStab","attackSlash","attackCrush","attackMagic","attackRanged",
    "defStab","defSlash","defCrush","defMagic","defRanged",
]

wb = Workbook()
ws = wb.active
ws.title = "Items"

header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
header_fill = PatternFill("solid", start_color="1F4E79")
cell_font = Font(name="Arial", size=10)
center = Alignment(horizontal="center")
thin = Side(style="thin", color="CCCCCC")
border = Border(bottom=thin)

# Freeze header row
ws.freeze_panes = "A2"

for col_idx, (label, width) in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=label)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[get_column_letter(col_idx)].width = width

ws.row_dimensions[1].height = 18

for row_idx, row in enumerate(rows, start=2):
    alt_fill = PatternFill("solid", start_color="EBF3FB") if row_idx % 2 == 0 else None
    for col_idx, field in enumerate(fields, start=1):
        val = row[field]
        if isinstance(val, bool):
            val = "Yes" if val else ""
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = cell_font
        if col_idx >= 8:
            cell.alignment = center
        if alt_fill:
            cell.fill = alt_fill

# Auto-filter on headers
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

out_path = "data/items/items-catalog.xlsx"
wb.save(out_path)
print(f"Saved {len(rows)} items to {out_path}")
